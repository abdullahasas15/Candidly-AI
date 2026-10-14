import asyncio
import os
import base64
from httpx import AsyncClient, ASGITransport
from docx import Document
from PIL import Image
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from backend.app.main import app
from backend.app.database import AsyncSessionLocal
from backend.app import models
from scripts.test_helpers import create_test_accounts

async def test_storage_integrity():
    print("=================================================================")
    print("TESTING CANDIDATE & RECRUITER DATABASE STORAGE & ARCHITECTURE")
    print("=================================================================")

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        recruiter_headers, candidate_headers, recruiter_user, candidate_user = await create_test_accounts(client, "storage")
        # -------------------------------------------------------------
        # STEP 1: CREATE FULL RECRUITER JOB POSTING WITH EXTENSIVE RUBRICS
        # -------------------------------------------------------------
        recruiter_job_payload = {
            "job_title": "Principal Distributed Systems Architect",
            "department": "Core Infrastructure",
            "seniority_level": "Senior (L5)",
            "employment_type": "Full-Time",
            "workplace_model": "Remote",
            "permitted_locations": ["Global Remote", "United States", "San Francisco, CA"],
            "total_experience_years": 8.0,
            "domain_experience_years": 6.0,
            "salary_range_min": 210000.0,
            "salary_range_max": 280000.0,
            "salary_currency": "USD",
            "skills": [
                {
                    "name": "Distributed Systems & Consensus (Raft/Paxos)",
                    "category": "Core Technical",
                    "priority_tier": "P0",
                    "weight_percentage": 40.0,
                    "rubric_l1": "Knows definitions of CAP theorem and state machines.",
                    "rubric_l2": "Has configured Raft/Paxos based tools like etcd or ZooKeeper.",
                    "rubric_l3": "Can design replicated log architectures and handle split-brain partitions.",
                    "rubric_l4": "Implements linearizable read paths and custom snapshotting / log compaction.",
                    "rubric_l5": "Formally models consensus edge cases with TLA+ and authors custom consensus protocols."
                },
                {
                    "name": "High-Throughput Stream Processing (Kafka/Flink)",
                    "category": "Core Technical",
                    "priority_tier": "P0",
                    "weight_percentage": 35.0,
                    "rubric_l1": "Basic producer/consumer concepts.",
                    "rubric_l2": "Configures consumer groups, offsets, and basic partitioning.",
                    "rubric_l3": "Guarantees exactly-once semantics and mitigates rebalancing storms.",
                    "rubric_l4": "Tunes low-level OS pagecache, zero-copy socket I/O, and JVM memory layouts.",
                    "rubric_l5": "Designs petabyte-scale multi-datacenter stream pipelines with sub-millisecond tail latency."
                },
                {
                    "name": "System Observability & Chaos Engineering",
                    "category": "Tooling",
                    "priority_tier": "P1",
                    "weight_percentage": 25.0,
                    "rubric_l1": "Looks at metrics dashboards.",
                    "rubric_l2": "Sets up Prometheus alerts and distributed tracing spans.",
                    "rubric_l3": "Conducts Chaos Engineering experiments with simulated network partitions.",
                    "rubric_l4": "Creates adaptive rate-limiting and circuit-breaking algorithms under brownout conditions.",
                    "rubric_l5": "Builds autonomous self-healing distributed control planes with provable safety invariants."
                }
            ],
            "certifications": [
                {
                    "name": "CKA (Certified Kubernetes Administrator)",
                    "issuing_org": "CNCF",
                    "priority_tier": "P1",
                    "verification_required": False
                },
                {
                    "name": "AWS Certified Solutions Architect - Professional",
                    "issuing_org": "Amazon Web Services",
                    "priority_tier": "P1",
                    "verification_required": False
                }
            ],
            "custom_questions": [
                {
                    "question_text": "Describe the most challenging split-brain partition or data corruption incident you resolved in production.",
                    "category": "Architecture",
                    "difficulty_level": "Principal",
                    "expected_key_points": ["Root cause analysis", "Network partition containment", "Data reconciliation"],
                    "priority_tier": "P0"
                }
            ]
        }

        res = await client.post("/api/jobs", json=recruiter_job_payload, headers=recruiter_headers)
        assert res.status_code == 201, f"Failed to create recruiter job posting: {res.text}"
        job = res.json()
        job_id = job["id"]
        print(f"✓ Recruiter Job Created: '{job['job_title']}' (ID: {job_id})")
        print(f"  - Skills configured: {len(job['skills'])}")
        print(f"  - Certifications configured: {len(job['certifications'])}")
        print(f"  - Custom Questions configured: {len(job['custom_questions'])}")
        print(f"  - Auto-Budgeted Question Plans: {len(job['question_plans'])}")

        # Verify weights sum to 100%
        total_weight = sum(s["weight_percentage"] for s in job["skills"])
        assert abs(total_weight - 100.0) < 0.001, f"Skill weights sum to {total_weight}, expected 100%"
        print(f"  - Total Skill Weight: {total_weight}% (Strict 100% calibration verified)")

        # -------------------------------------------------------------
        # STEP 2: CREATE CANDIDATE APPLICATION & STORE COMPREHENSIVE DOSSIER
        # -------------------------------------------------------------
        candidate_app_payload = {
            "job_id": job_id,
            "candidate_id": candidate_user["id"],
            "full_name": "Elena Rostova",
            "email": candidate_user["email"],
            "phone": "+1 (415) 890-1234",
            "linkedin_url": "https://linkedin.com/in/elena-rostova-dist",
            "portfolio_url": "https://github.com/elena-dist-systems",
            "self_reported_experience_years": 9.0,
            "fit_pitch": "Lead architect of open-source Raft consensus engine used by 40+ fintech enterprises.",
            "expected_salary": 245000.0,
            "expected_salary_currency": "USD"
        }
        res = await client.post("/api/applications", json=candidate_app_payload, headers=candidate_headers)
        assert res.status_code == 201, f"Failed to create application draft: {res.text}"
        app_record = res.json()
        app_id = app_record["id"]
        print(f"\n✓ Candidate Application Created: ID {app_id} ({app_record['full_name']})")

        # -------------------------------------------------------------
        # STEP 3: UPLOAD REAL DOCX RESUME & RUN PARSER
        # -------------------------------------------------------------
        docx_path = "scripts/test_elena_resume.docx"
        doc = Document()
        doc.add_heading("Elena Rostova - Distributed Systems Architect", level=1)
        doc.add_paragraph("Summary: Experienced architect specializing in Raft, Paxos, Kafka, Go, and C++.")
        doc.add_paragraph("Work Experience:")
        doc.add_paragraph("Datastream Core - Principal Architect (Jan 2020 - Present)")
        doc.add_paragraph("Designed zero-copy log replication engine with Raft consensus.")
        doc.add_paragraph("HyperScale Systems - Senior Distributed Engineer (Jan 2015 - Dec 2019)")
        doc.add_paragraph("Built high-frequency order book pipeline processing 2M events/sec.")
        doc.add_paragraph("Education:")
        doc.add_paragraph("M.S. in Computer Science - Carnegie Mellon University (2014)")
        doc.add_paragraph("B.S. in Computer Science & Applied Mathematics - UC Berkeley (2012)")
        doc.add_paragraph("Certifications:")
        doc.add_paragraph("AWS Certified Solutions Architect - Professional (2022)")
        doc.add_paragraph("Certified Kubernetes Administrator (CKA) (2021)")
        doc.add_paragraph("Skills:")
        doc.add_paragraph("Go, C++, Rust, Python, Raft, Paxos, Kafka, Flink, Kubernetes, Docker, PostgreSQL, Redis, Chaos Engineering")
        doc.save(docx_path)

        with open(docx_path, "rb") as f:
            files = {"file": ("test_elena_resume.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
            res = await client.post(f"/api/applications/{app_id}/resume", files=files, headers=candidate_headers)
        assert res.status_code == 200, f"Failed to ingest resume: {res.text}"
        resume_res = res.json()
        parsed = resume_res["parsed_structured_data"]
        print(f"✓ Resume Ingested & Parsed Successfully:")
        print(f"  - Extracted Text Length: {len(resume_res['parsed_text'])} characters")
        print(f"  - Total Computed Experience: {parsed['total_experience_years']:.1f} years")
        print(f"  - Categorized Skills: {list(parsed.get('skills_by_category', {}).keys())}")
        print(f"  - Education Entries: {len(parsed['education'])}")
        print(f"  - Certifications Entries: {len(parsed['certifications'])}")
        print(f"  - Projects Entries: {len(parsed.get('projects', []))}")

        # -------------------------------------------------------------
        # STEP 4: EDIT & SAVE HIGH-FIDELITY CANDIDATE DOSSIER
        # -------------------------------------------------------------
        enriched_structured_data = {
            **parsed,
            "categorized_skills": {
                "languages": ["Go", "C++", "Rust", "Python"],
                "frameworks": ["gRPC", "Protobuf", "Apache Flink"],
                "databases": ["PostgreSQL", "RocksDB", "CockroachDB", "Redis"],
                "cloud_devops": ["Kubernetes", "AWS", "Docker", "Terraform"],
                "tools_methods": ["Chaos Mesh", "TLA+", "Prometheus", "Grafana", "Jepsen Testing"]
            },
            "education": [
                {
                    "degree": "Master of Science in Computer Science",
                    "institution": "Carnegie Mellon University",
                    "field_of_study": "Distributed Systems",
                    "graduation_year": "2014",
                    "gpa": "3.95"
                },
                {
                    "degree": "Bachelor of Science",
                    "institution": "UC Berkeley",
                    "field_of_study": "Computer Science & Applied Mathematics",
                    "graduation_year": "2012",
                    "gpa": "3.90"
                }
            ],
            "certifications": [
                {
                    "name": "AWS Certified Solutions Architect - Professional",
                    "issuer": "Amazon Web Services",
                    "issue_year": "2022",
                    "credential_id": "AWS-PSA-994821"
                },
                {
                    "name": "CKA (Certified Kubernetes Administrator)",
                    "issuer": "CNCF",
                    "issue_year": "2021",
                    "credential_id": "CKA-77310"
                }
            ],
            "projects": [
                {
                    "title": "AegisRaft: Low-Latency Raft Implementation in Rust",
                    "role": "Lead Architect & Maintainer",
                    "description": "Engineered pipelined Raft log replication engine achieving 450k op/s with zero-allocation buffers.",
                    "tech_stack": ["Rust", "gRPC", "RocksDB", "TLA+"],
                    "outcomes": "Adopted by 3 production fintech backends; Jepsen tested with 0 linearizability violations."
                },
                {
                    "title": "ChaosStream: Automated Network Partition Simulator",
                    "role": "Creator",
                    "description": "eBPF-driven network latency and packet loss injector for multi-AZ Kubernetes clusters.",
                    "tech_stack": ["Go", "eBPF", "Kubernetes CRDs"],
                    "outcomes": "Detected 14 silent split-brain edge cases prior to production deployments."
                }
            ]
        }

        patch_res = await client.patch(
            f"/api/applications/{app_id}/parsed-resume",
            json={
                "parsed_structured_data": enriched_structured_data,
                "self_reported_experience_years": 9.0
            },
            headers=candidate_headers,
        )
        assert patch_res.status_code == 200, f"Failed to patch parsed resume: {patch_res.text}"
        print("✓ Enriched Structured Candidate Dossier Persisted via API")

        # -------------------------------------------------------------
        # STEP 5: STORE WEBCAM BIOMETRIC BASELINE
        # -------------------------------------------------------------
        img = Image.new("RGB", (320, 240), color=(30, 80, 160))
        test_cam_path = "scripts/test_elena_webcam.jpg"
        img.save(test_cam_path, "JPEG")
        with open(test_cam_path, "rb") as f:
            b64_str = "data:image/jpeg;base64," + base64.b64encode(f.read()).decode("utf-8")

        verif_res = await client.post(
            f"/api/applications/{app_id}/verification",
            json={"live_photo": b64_str, "government_id": None},
            headers=candidate_headers,
        )
        assert verif_res.status_code == 200
        print("✓ Webcam Biometric Baseline Photo Stored & Verified")

        # -------------------------------------------------------------
        # STEP 6: SUBMIT APPLICATION
        # -------------------------------------------------------------
        submit_res = await client.patch(
            f"/api/applications/{app_id}/submit",
            json={"fit_pitch": "Extensive 9-year record architecting fault-tolerant consensus systems and low-latency pipelines."},
            headers=candidate_headers,
        )
        assert submit_res.status_code == 200
        print("✓ Candidate Application Finalized & Submitted")

    # -------------------------------------------------------------
    # STEP 7: DIRECT POSTGRESQL ORM VERIFICATION OF STORAGE ARCHITECTURE
    # -------------------------------------------------------------
    print("\n--- Verifying Raw PostgreSQL Database Records ---")
    async with AsyncSessionLocal() as db:
        # Check Job Record
        result = await db.execute(
            select(models.JobPosting)
            .options(
                selectinload(models.JobPosting.skills),
                selectinload(models.JobPosting.certifications),
                selectinload(models.JobPosting.custom_questions),
                selectinload(models.JobPosting.question_plans),
            )
            .filter(models.JobPosting.id == job_id)
        )
        db_job = result.scalars().first()
        assert db_job is not None, "Job record missing from PostgreSQL"
        print(f"✓ DB Job Posting: '{db_job.job_title}'")
        print(f"  - Salary: ${db_job.salary_range_min:,.0f} - ${db_job.salary_range_max:,.0f} {db_job.salary_currency}")
        print(f"  - Skills in DB: {len(db_job.skills)}")
        for s in db_job.skills:
            assert s.rubric_l1 is not None and s.rubric_l5 is not None
            print(f"    * {s.name} [{s.priority_tier}] (Weight: {s.weight_percentage}%): L3='{s.rubric_l3[:30]}...', L5='{s.rubric_l5[:30]}...'")
        
        print(f"  - Certifications in DB: {len(db_job.certifications)}")
        print(f"  - Custom Questions in DB: {len(db_job.custom_questions)}")
        print(f"  - Question Plans in DB: {len(db_job.question_plans)}")

        # Check Application Record
        result = await db.execute(
            select(models.JobApplication)
            .options(
                selectinload(models.JobApplication.resume),
                selectinload(models.JobApplication.id_verification),
                selectinload(models.JobApplication.discrepancy_flags),
            )
            .filter(models.JobApplication.id == app_id)
        )
        db_app = result.scalars().first()
        assert db_app is not None, "Application record missing from PostgreSQL"
        print(f"\n✓ DB Application: '{db_app.full_name}' ({db_app.email})")
        print(f"  - Status: '{db_app.status}'")
        print(f"  - Expected Salary: ${db_app.expected_salary:,.0f} {db_app.expected_salary_currency}")
        print(f"  - Fit Pitch: '{db_app.fit_pitch}'")

        # Check Resume & Structured JSON Fields
        db_resume = db_app.resume
        assert db_resume is not None, "Resume record missing from PostgreSQL"
        struct = db_resume.parsed_structured_data
        assert "categorized_skills" in struct, "categorized_skills missing from parsed_structured_data"
        assert "education" in struct and len(struct["education"]) == 2, "Education JSON entries incorrect"
        assert "certifications" in struct and len(struct["certifications"]) == 2, "Certifications JSON entries incorrect"
        assert "projects" in struct and len(struct["projects"]) == 2, "Projects JSON entries incorrect"

        print(f"\n✓ DB Candidate Structured Dossier (JSON Storage Integrity):")
        print(f"  - Languages: {struct['categorized_skills']['languages']}")
        print(f"  - Databases: {struct['categorized_skills']['databases']}")
        print(f"  - Cloud/DevOps: {struct['categorized_skills']['cloud_devops']}")
        print(f"  - Tools/Methods: {struct['categorized_skills']['tools_methods']}")
        print(f"  - Verified Education Count: {len(struct['education'])}")
        for edu in struct["education"]:
            print(f"    * {edu['degree']} at {edu['institution']} ({edu['graduation_year']})")
        print(f"  - Verified Certifications Count: {len(struct['certifications'])}")
        for cert in struct["certifications"]:
            print(f"    * {cert['name']} ({cert['issuer']})")
        print(f"  - Verified Projects Count: {len(struct['projects'])}")
        for proj in struct["projects"]:
            print(f"    * {proj['title']} (Role: {proj['role']})")

        # Check Biometric ID Verification
        db_verif = db_app.id_verification
        assert db_verif is not None, "ID Verification missing from PostgreSQL"
        if db_verif.live_photo_path.startswith("data:image/"):
            assert ";base64," in db_verif.live_photo_path
            storage_description = "data URI"
        else:
            disk_path = db_verif.live_photo_path.lstrip('/')
            assert os.path.exists(disk_path), f"Webcam file missing on disk at {disk_path}"
            storage_description = "file on disk"
        print(f"\n✓ DB Biometric Record:")
        print(f"  - Photo Path: {db_verif.live_photo_path[:80]}... (Storage: {storage_description})")
        print(f"  - Status: '{db_verif.verification_status}'")

    print("\n=================================================================")
    print("ALL CANDIDATE & RECRUITER DATA STORAGE TESTS PASSED SUCCESSFULLY!")
    print("Database is 100% prepared for custom LLM question generation.")
    print("=================================================================")

if __name__ == "__main__":
    asyncio.run(test_storage_integrity())
